import csv
import io
import json
import math
import os
import secrets
import sqlite3
import uuid
from datetime import date, datetime, timedelta
from zipfile import BadZipFile
from flask import Flask, abort, flash, redirect, render_template, request, url_for
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import pandas as pd
import numpy as np
import pickle
import re
from pathlib import Path
from project_store import connect, delete_project, get_project, get_projects, init_db, save_project

app = Flask(__name__)
app.secret_key = os.environ.get("SESSION_SECRET") or secrets.token_hex(32)
app.config["MAX_CONTENT_LENGTH"] = 3 * 1024 * 1024

# ============================================================
# FILE PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent
DATABASE_FILE = BASE_DIR / "data" / "uptrack_projects.sqlite3"
app.config["DATABASE"] = str(DATABASE_FILE)
init_db(DATABASE_FILE)

APRIL_FILE = BASE_DIR / "data" / "april_projects.csv"
MAY_FILE = BASE_DIR / "data" / "may_projects.csv"
JUNE_FILE = BASE_DIR / "data" / "june_projects.csv"
ML_DATASET_FILE = BASE_DIR / "data" / "uptrack_ml_dataset.csv"

MODEL_FILE = BASE_DIR / "models" / "uptrack_model_balanced.pkl"
SCALER_FILE = BASE_DIR / "models" / "uptrack_scaler.pkl"


# ============================================================
# LOAD DATA
# ============================================================

df = pd.read_csv(APRIL_FILE)
may_df = pd.read_csv(MAY_FILE)
june_df = pd.read_csv(JUNE_FILE)
ml_dataset_df = pd.read_csv(ML_DATASET_FILE)


# ============================================================
# LOAD MODEL
# ============================================================

with open(MODEL_FILE, "rb") as file:
    model = pickle.load(file)

with open(SCALER_FILE, "rb") as file:
    scaler = pickle.load(file)


# ============================================================
# MODEL FEATURES
# ============================================================

FEATURE_COLUMNS = [
    "april_anticipated_month",
    "april_original_commissioning_month",
    "april_revised_commissioning_month",
    "april_original_cost_cr",
    "april_revised_cost_cr",
    "april_anticipated_cost_cr",
    "april_cumulative_expenditure_cr",
    "april_cost_overrun_cr",
    "april_time_overrun_months"
]

REFERENCE = pd.Period("2000-01", freq="M")


# ============================================================
# PROJECT NAME NORMALIZATION
# ============================================================

def normalize_project_name(value):

    if pd.isna(value):
        return ""

    value = str(value).upper().strip()

    value = re.sub(r"[^A-Z0-9]+", " ", value)

    value = re.sub(r"\s+", " ", value)

    return value.strip()


df["_normalized_project_name"] = df["project_name"].apply(
    normalize_project_name
)

may_df["_normalized_project_name"] = may_df["project_name"].apply(
    normalize_project_name
)

june_df["_normalized_project_name"] = june_df["project_name"].apply(
    normalize_project_name
)


# ============================================================
# FIND PROJECT IN MONTHLY DATA
# ============================================================

def find_month_project(normalized_name, monthly_df):

    if not normalized_name:
        return None

    matches = monthly_df[
        monthly_df["_normalized_project_name"] == normalized_name
    ]

    if len(matches) == 1:

        row = matches.iloc[0].copy()

        row = row.astype(object).where(
            pd.notna(row),
            None
        )

        return row

    return None


# ============================================================
# CLEAN VALUE
# ============================================================

def clean_value(value):

    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if str(value).strip().lower() in [
        "",
        "nan",
        "none"
    ]:
        return None

    return value


# ============================================================
# MONTH NUMBER
# ============================================================

def month_number(value):

    value = clean_value(value)

    if value is None:
        return None

    try:

        parsed = pd.to_datetime(
            str(value),
            errors="coerce"
        )

        if pd.isna(parsed):
            return None

        period = parsed.to_period("M")

        return (
            (period.year - REFERENCE.year) * 12
            + (period.month - REFERENCE.month)
        )

    except Exception:

        return None


# ============================================================
# NUMERIC VALUE
# ============================================================

def numeric_value(value):

    value = clean_value(value)

    if value is None:
        return None

    try:

        return float(
            str(value).replace(",", "").strip()
        )

    except (ValueError, TypeError):

        return None


# ============================================================
# FEATURE MEDIANS
# ============================================================

# ============================================================
# CREATE MODEL FEATURES
# ============================================================

def make_features(project):

    features = {}

    features["april_anticipated_month"] = month_number(
        project.get("anticipated_date_of_commissioning")
    )

    features["april_original_commissioning_month"] = month_number(
        project.get("original_date_of_commissioning")
    )

    features["april_revised_commissioning_month"] = month_number(
        project.get("revised_date_of_commissioning")
    )

    features["april_original_cost_cr"] = numeric_value(
        project.get("original_cost_cr")
    )

    features["april_revised_cost_cr"] = numeric_value(
        project.get("revised_cost_cr")
    )

    features["april_anticipated_cost_cr"] = numeric_value(
        project.get("anticipated_cost_cr")
    )

    features["april_cumulative_expenditure_cr"] = numeric_value(
        project.get("cumulative_expenditure_cr")
    )

    features["april_cost_overrun_cr"] = numeric_value(
        project.get("cost_overrun_cr")
    )

    features["april_time_overrun_months"] = numeric_value(
        project.get("time_overrun_months")
    )

    missing = [column for column in FEATURE_COLUMNS if features.get(column) is None or pd.isna(features.get(column))]
    if missing:
        raise ValueError(
            "Required model inputs are missing or invalid: "
            + ", ".join(missing)
            + ". No values were estimated."
        )

    values = [float(features[column]) for column in FEATURE_COLUMNS]
    return np.array(values).reshape(1, -1)


# ============================================================
# PREDICTION
# ============================================================

def predict_delay(project):

    features = make_features(project)

    scaled_features = scaler.transform(features)

    prediction = int(
        model.predict(scaled_features)[0]
    )

    probability = float(
        model.predict_proba(
            scaled_features
        )[0][1]
    )

    return prediction, probability


# ============================================================
# NUMERIC CHANGE COMPARISON
# ============================================================

def compare_numeric(
    april_value,
    may_value,
    june_value,
    label,
    unit=""
):

    values = [
        ("April", clean_value(april_value)),
        ("May", clean_value(may_value)),
        ("June", clean_value(june_value))
    ]

    available = []

    for month, value in values:

        if value is not None:

            try:

                numeric = float(
                    str(value).replace(",", "")
                )

                available.append(
                    (month, numeric)
                )

            except (ValueError, TypeError):
                pass

    if len(available) < 2:
        return None

    first_month, first_value = available[0]
    last_month, last_value = available[-1]

    if first_value == last_value:

        return (
            f"✓ {label} remained unchanged: "
            f"{first_value:g}{unit}"
        )

    if last_value > first_value:

        return (
            f"⚠ {label} increased: "
            f"{first_value:g}{unit} → "
            f"{last_value:g}{unit}"
        )

    return (
        f"✓ {label} decreased: "
        f"{first_value:g}{unit} → "
        f"{last_value:g}{unit}"
    )


# ============================================================
# COMPLETION DATE COMPARISON
# ============================================================

def compare_completion_date(
    april_value,
    may_value,
    june_value
):

    def parse_project_month(value):

        value = clean_value(value)

        if value is None:
            return None

        text = str(value).strip()

        match = re.fullmatch(
            r"(\d{1,2})/(\d{4})",
            text
        )

        if match:

            month = int(match.group(1))
            year = int(match.group(2))

            if 1 <= month <= 12:

                return pd.Period(
                    year=year,
                    month=month,
                    freq="M"
                )

            return None

        match = re.fullmatch(
            r"([A-Za-z]{3})-(\d{2,4})",
            text
        )

        if match:

            month_text = match.group(1).title()
            year_text = match.group(2)

            month_map = {
                "Jan": 1,
                "Feb": 2,
                "Mar": 3,
                "Apr": 4,
                "May": 5,
                "Jun": 6,
                "Jul": 7,
                "Aug": 8,
                "Sep": 9,
                "Oct": 10,
                "Nov": 11,
                "Dec": 12
            }

            month = month_map.get(month_text)

            if month is None:
                return None

            year = int(year_text)

            if len(year_text) == 2:
                year += 2000

            return pd.Period(
                year=year,
                month=month,
                freq="M"
            )

        parsed = pd.to_datetime(
            text,
            errors="coerce"
        )

        if pd.isna(parsed):
            return None

        return pd.Period(
            year=parsed.year,
            month=parsed.month,
            freq="M"
        )

    values = [
        ("April", april_value),
        ("May", may_value),
        ("June", june_value)
    ]

    available = []

    for month_name, value in values:

        parsed = parse_project_month(value)

        if parsed is not None:

            available.append(
                (month_name, parsed)
            )

    if len(available) < 2:
        return None

    first_month, first_date = available[0]
    last_month, last_date = available[-1]

    first_text = first_date.strftime("%b-%y")
    last_text = last_date.strftime("%b-%y")

    if first_date == last_date:

        return (
            f"✓ Completion date remained unchanged: "
            f"{first_text}"
        )

    if last_date > first_date:

        return (
            f"⚠ Completion date moved later: "
            f"{first_text} → {last_text}"
        )

    return (
        f"✓ Completion date moved earlier: "
        f"{first_text} → {last_text}"
    )


# ============================================================
# BUILD WHAT CHANGED
# ============================================================

def build_changes(
    project,
    may_project,
    june_project
):

    changes = []

    if project is None:
        return changes

    april_completion = project.get(
        "anticipated_date_of_commissioning"
    )

    april_revised_cost = project.get(
        "revised_cost_cr"
    )

    april_expenditure = project.get(
        "cumulative_expenditure_cr"
    )

    april_time_overrun = project.get(
        "time_overrun_months"
    )

    may_completion = None
    may_revised_cost = None
    may_expenditure = None
    may_time_overrun = None

    if may_project is not None:

        may_completion = may_project.get(
            "anticipated_date_of_commissioning"
        )

        may_revised_cost = may_project.get(
            "revised_cost_cr"
        )

        may_expenditure = may_project.get(
            "cumulative_expenditure_cr"
        )

        may_time_overrun = may_project.get(
            "time_overrun_months"
        )

    june_completion = None
    june_revised_cost = None
    june_expenditure = None
    june_time_overrun = None

    if june_project is not None:

        june_completion = june_project.get(
            "anticipated_date_of_commissioning"
        )

        june_revised_cost = june_project.get(
            "revised_cost_cr"
        )

        june_expenditure = june_project.get(
            "cumulative_expenditure_cr"
        )

        june_time_overrun = june_project.get(
            "time_overrun_months"
        )

    completion_change = compare_completion_date(
        april_completion,
        may_completion,
        june_completion
    )

    if completion_change:
        changes.append(completion_change)

    revised_cost_change = compare_numeric(
        april_revised_cost,
        may_revised_cost,
        june_revised_cost,
        "Revised cost",
        " Cr"
    )

    if revised_cost_change:
        changes.append(revised_cost_change)

    expenditure_change = compare_numeric(
        april_expenditure,
        may_expenditure,
        june_expenditure,
        "Cumulative expenditure",
        " Cr"
    )

    if expenditure_change:
        changes.append(expenditure_change)

    time_overrun_change = compare_numeric(
        april_time_overrun,
        may_time_overrun,
        june_time_overrun,
        "Time overrun",
        " months"
    )

    if time_overrun_change:
        changes.append(time_overrun_change)

    return changes


# ============================================================
# MAIN PAGE
# ============================================================

@app.route("/research", methods=["GET", "POST"])
def research():

    project_names = (
        df["project_name"]
        .dropna()
        .astype(str)
        .tolist()
    )

    total_april_projects = len(df)

    ml_dataset_projects = len(ml_dataset_df)

    delay_projects = int(
        (ml_dataset_df["delay"] == 1).sum()
    )

    no_delay_projects = int(
        (ml_dataset_df["delay"] == 0).sum()
    )

    selected_name = ""

    project = None
    may_project = None
    june_project = None

    prediction = None
    probability = None
    prediction_message = None

    changes = []

    # False unless Predict Delay is clicked
    show_prediction = False

    # ========================================================
    # WHAT-IF VARIABLES
    # ========================================================

    simulation_prediction = None
    simulation_probability = None
    simulation_date = ""
    simulation_message = None

    if request.method == "POST":

        selected_name = request.form.get(
            "project_name",
            ""
        )

        action = request.form.get(
            "action",
            "view"
        )

        if selected_name:

            selected_rows = df[
                df["project_name"].astype(str)
                == selected_name
            ]

            if not selected_rows.empty:

                project = selected_rows.iloc[0]

                normalized_name = normalize_project_name(
                    selected_name
                )

                may_project = find_month_project(
                    normalized_name,
                    may_df
                )

                june_project = find_month_project(
                    normalized_name,
                    june_df
                )

                changes = build_changes(
                    project,
                    may_project,
                    june_project
                )

                # ====================================================
                # SHOW PROJECT
                # ====================================================

                if action == "view":

                    prediction = None
                    probability = None
                    show_prediction = False

                # ====================================================
                # PREDICT DELAY
                # ====================================================

                elif action == "predict":

                    show_prediction = True

                    try:

                        prediction, probability = predict_delay(
                            project
                        )

                    except Exception as error:

                        print(
                            "Prediction error:",
                            error
                        )

                        prediction = None
                        probability = None
                        prediction_message = (
                            str(error) if isinstance(error, ValueError)
                            else "The experimental model could not evaluate this record."
                        )

                # ====================================================
                # WHAT-IF SIMULATION
                # ====================================================

                elif action == "simulate":

                    simulation_date = request.form.get(
                        "simulation_date",
                        ""
                    )

                    if simulation_date:

                        try:

                            # Make a temporary copy.
                            # Original project data is NOT changed.
                            simulated_project = project.copy()

                            # Change only the expected
                            # completion date.
                            simulated_project[
                                "anticipated_date_of_commissioning"
                            ] = simulation_date

                            (
                                simulation_prediction,
                                simulation_probability
                            ) = predict_delay(
                                simulated_project
                            )

                        except Exception as error:

                            print(
                                "Simulation error:",
                                error
                            )

                            simulation_prediction = None
                            simulation_probability = None
                            simulation_message = (
                                str(error) if isinstance(error, ValueError)
                                else "The experimental model could not evaluate this scenario."
                            )

    return render_template(
        "index.html",

        total_april_projects=total_april_projects,

        ml_dataset_projects=ml_dataset_projects,

        delay_projects=delay_projects,

        no_delay_projects=no_delay_projects,

        project_names=project_names,

        selected_name=selected_name,

        project=project,

        may_project=may_project,

        june_project=june_project,

        prediction=prediction,

        probability=probability,

        show_prediction=show_prediction,

        changes=changes,

        # What-If values
        simulation_prediction=simulation_prediction,

        simulation_probability=simulation_probability,

        simulation_date=simulation_date,
        prediction_message=prediction_message,
        simulation_message=simulation_message
    )


PROJECT_STATUSES = ("Planned", "In Progress", "On Hold", "Completed")
IMPORT_LIMIT = 1000
IMPORT_COLUMNS = {
    "project_id": ("project_id", "id", "unique_id"),
    "name": ("name", "project_name", "project"),
    "project_type": ("project_type", "type"),
    "location": ("location",),
    "start_date": ("start_date", "project_start_date"),
    "original_completion_date": ("original_completion_date", "original_end_date"),
    "expected_completion_date": ("expected_completion_date", "current_expected_completion_date", "revised_completion_date"),
    "budget": ("budget", "original_budget"),
    "expenditure": ("expenditure", "spent", "cumulative_expenditure"),
    "progress": ("progress", "progress_percent", "progress_percentage"),
    "status": ("status", "project_status"),
    "milestones": ("milestones", "milestone"),
    "notes": ("notes", "note"),
}


def _date_value(value, label, errors):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    try:
        return date.fromisoformat(str(value).strip()[:10]).isoformat()
    except ValueError:
        errors.append(f"{label} must be a valid date in YYYY-MM-DD format.")
        return None


def _number_value(value, label, errors, *, required=False, maximum=None):
    if value is None or str(value).strip() == "":
        if required:
            errors.append(f"{label} is required.")
        return 0.0 if required else None
    try:
        number = float(str(value).replace(",", "").replace("₹", "").strip())
    except (ValueError, TypeError):
        errors.append(f"{label} must be a number.")
        return None
    if not math.isfinite(number) or number < 0:
        errors.append(f"{label} must be zero or greater.")
        return None
    if maximum is not None and number > maximum:
        errors.append(f"{label} must be no more than {maximum:g}.")
        return None
    return number


def _parse_milestones(value, errors):
    if isinstance(value, list):
        raw_lines = [str(item.get("name", "")) + (" | " + str(item.get("due_date")) if item.get("due_date") else "") for item in value if isinstance(item, dict)]
    else:
        raw_lines = str(value or "").splitlines()
    milestones = []
    for line in raw_lines:
        line = line.strip()
        if not line:
            continue
        pieces = [piece.strip() for piece in line.split("|", 1)]
        name = pieces[0]
        if not name:
            errors.append("Milestone names cannot be blank.")
            continue
        if len(name) > 120:
            errors.append("Milestone names must be 120 characters or fewer.")
            continue
        due = _date_value(pieces[1], "Milestone date", errors) if len(pieces) == 2 and pieces[1] else None
        milestones.append({"name": name, "due_date": due, "complete": False})
    if len(milestones) > 30:
        errors.append("A project can have at most 30 milestones.")
    return milestones[:30]


def _validate_project(raw):
    errors = []
    project_id = str(raw.get("project_id") or "").strip()
    name = str(raw.get("name") or "").strip()
    if not project_id:
        errors.append("Project ID is required.")
    elif len(project_id) > 60:
        errors.append("Project ID must be 60 characters or fewer.")
    if not name:
        errors.append("Project name is required.")
    elif len(name) > 160:
        errors.append("Project name must be 160 characters or fewer.")

    start = _date_value(raw.get("start_date"), "Start date", errors)
    original = _date_value(raw.get("original_completion_date"), "Original completion date", errors)
    expected = _date_value(raw.get("expected_completion_date"), "Expected completion date", errors)
    for label, value in (("Original completion date", original), ("Expected completion date", expected)):
        if start and value and value < start:
            errors.append(f"{label} cannot be before the start date.")

    budget = _number_value(raw.get("budget"), "Budget", errors)
    expenditure = _number_value(raw.get("expenditure"), "Expenditure", errors)
    progress = _number_value(raw.get("progress"), "Progress", errors, maximum=100)
    if progress is None:
        progress = 0.0
    status = str(raw.get("status") or "Planned").strip()
    if status not in PROJECT_STATUSES:
        errors.append("Status must be Planned, In Progress, On Hold, or Completed.")
    milestones = _parse_milestones(raw.get("milestones"), errors)
    project_type = str(raw.get("project_type") or "").strip()
    location = str(raw.get("location") or "").strip()
    notes = str(raw.get("notes") or "").strip()
    if len(project_type) > 100:
        errors.append("Project type must be 100 characters or fewer.")
    if len(location) > 160:
        errors.append("Location must be 160 characters or fewer.")
    if len(notes) > 3000:
        errors.append("Notes must be 3,000 characters or fewer.")
    return {
        "project_id": project_id,
        "name": name,
        "project_type": project_type,
        "location": location,
        "start_date": start,
        "original_completion_date": original,
        "expected_completion_date": expected,
        "budget": budget,
        "expenditure": expenditure,
        "progress": progress,
        "status": status,
        "milestones": milestones,
        "notes": notes,
    }, errors


def _read_import(file_storage):
    filename = (file_storage.filename or "").lower()
    if not filename.endswith((".csv", ".xlsx")):
        raise ValueError("Choose a CSV or Excel .xlsx file.")
    payload = file_storage.read()
    if not payload:
        raise ValueError("The selected file is empty.")
    try:
        if filename.endswith(".csv"):
            text = payload.decode("utf-8-sig")
            source_rows = list(csv.DictReader(io.StringIO(text)))
        else:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                raise ValueError("The selected spreadsheet has no header row.")
            header_names = [str(value or "").strip() for value in headers]
            source_rows = [
                dict(zip(header_names, row))
                for row in values
                if any(value is not None and str(value).strip() for value in row)
            ]
            workbook.close()
    except UnicodeDecodeError as error:
        raise ValueError("CSV files must use UTF-8 text encoding.") from error
    except (csv.Error, OSError, BadZipFile, InvalidFileException) as error:
        raise ValueError("The file could not be read. Check that it is a valid CSV or .xlsx file.") from error
    if len(source_rows) > IMPORT_LIMIT:
        raise ValueError(f"Imports are limited to {IMPORT_LIMIT} rows per file.")
    if not source_rows:
        raise ValueError("No project rows were found in the file.")

    lookup = {}
    for key, aliases in IMPORT_COLUMNS.items():
        for alias in aliases:
            lookup[re.sub(r"[^a-z0-9]+", "_", alias.lower()).strip("_")] = key
    normalized_rows = []
    for source in source_rows:
        normalized = {}
        for key, value in source.items():
            normalized_key = re.sub(r"[^a-z0-9]+", "_", str(key).strip().lower()).strip("_")
            if normalized_key in lookup:
                normalized[lookup[normalized_key]] = value
        normalized_rows.append(normalized)
    return normalized_rows


def _project_warnings(project, today):
    warnings = []
    if project["status"] != "Completed" and project.get("expected_completion_date"):
        expected = date.fromisoformat(project["expected_completion_date"])
        if expected < today:
            warnings.append("Expected completion date has passed.")
        elif (expected - today).days <= 30:
            warnings.append("Expected completion is within 30 days.")
    if project["status"] == "In Progress" and not project.get("milestones"):
        warnings.append("No milestones have been recorded.")
    budget = project.get("budget")
    expenditure = project.get("expenditure")
    if budget is not None and expenditure is not None and expenditure > budget:
        warnings.append("Recorded expenditure is above the budget.")
    elif budget and expenditure is not None and project.get("progress", 0) < 100:
        spend_share = expenditure / budget
        progress_share = project.get("progress", 0) / 100
        if spend_share > progress_share + 0.2:
            warnings.append("Recorded spending is more than 20 percentage points ahead of progress.")
    return warnings


def _project_delay_assessment(project, today):
    """Explain schedule and budget signals without applying the historical ML model."""
    if project["status"] == "Completed":
        return {
            "level": "Completed",
            "tone": "complete",
            "reasons": ["This project is marked complete; no future delay risk is assessed."],
        }

    high_signals = []
    watch_signals = []
    context_signals = []
    reasons = []
    progress = float(project.get("progress") or 0)
    expected_raw = project.get("expected_completion_date")
    original_raw = project.get("original_completion_date")
    expected = date.fromisoformat(expected_raw) if expected_raw else None
    original = date.fromisoformat(original_raw) if original_raw else None

    if expected:
        days_left = (expected - today).days
        if days_left < 0:
            high_signals.append(f"The expected completion date passed {-days_left} day(s) ago.")
        elif days_left <= 14 and progress < 80:
            high_signals.append(
                f"The deadline is {days_left} day(s) away and recorded progress is {progress:g}%."
            )
        elif days_left <= 30 and progress < 90:
            watch_signals.append(
                f"The deadline is within {days_left} day(s), with {progress:g}% progress recorded."
            )
        else:
            context_signals.append("A current expected completion date is recorded.")
    else:
        reasons.append("No current expected completion date is recorded, so schedule timing is hard to assess.")

    if expected and original:
        slip_days = (expected - original).days
        if slip_days > 0:
            signal = f"The current expected date is {slip_days} day(s) later than the original date."
            (high_signals if slip_days > 30 else watch_signals).append(signal)
        else:
            context_signals.append("The current expected date is not later than the original date.")
    elif original:
        context_signals.append("An original completion date is recorded, but there is no current date to compare.")

    budget = project.get("budget")
    expenditure = project.get("expenditure")
    if budget is not None and expenditure is not None:
        if expenditure > budget:
            high_signals.append("Recorded expenditure is above the recorded budget.")
        elif budget > 0 and progress < 100:
            spend_gap = expenditure / budget - progress / 100
            if spend_gap >= 0.4:
                high_signals.append(
                    "Recorded spending is at least 40 percentage points ahead of progress."
                )
            elif spend_gap > 0.2:
                watch_signals.append(
                    "Recorded spending is more than 20 percentage points ahead of progress."
                )
            else:
                context_signals.append("Recorded spending is not substantially ahead of progress.")
    else:
        reasons.append("Budget risk cannot be compared until both budget and expenditure are recorded.")

    if project["status"] == "In Progress" and not project.get("milestones"):
        watch_signals.append("No milestones are recorded to track near-term delivery.")
    else:
        overdue_milestones = [
            milestone for milestone in project.get("milestones", [])
            if not milestone.get("complete")
            and milestone.get("due_date")
            and milestone["due_date"] < today.isoformat()
        ]
        for milestone in overdue_milestones:
            watch_signals.append(
                f"Milestone “{milestone['name']}” was due {milestone['due_date']}."
            )

    reasons = high_signals + watch_signals + reasons + context_signals
    if high_signals:
        level, tone = "High attention", "high"
    elif watch_signals:
        level, tone = "Watch", "watch"
    elif expected:
        level, tone = "Lower current concern", "low"
        reasons.insert(0, "No high-priority schedule or budget flags were found in the recorded details.")
    else:
        level, tone = "Need more data", "unknown"
        reasons.insert(0, "There is not enough schedule information to assess delay risk.")

    return {"level": level, "tone": tone, "reasons": reasons}


def _simulate_project(project, form, today):
    delay_text = str(form.get("delay_days") or "").strip()
    if delay_text:
        try:
            delay_days = int(delay_text)
        except ValueError as error:
            raise ValueError("Extra delay must be a whole number of days.") from error
    else:
        delay_days = 0
    if delay_days < 0 or delay_days > 3650:
        raise ValueError("Extra delay must be between 0 and 3,650 days.")

    changed_deadline = None
    changed_text = str(form.get("changed_deadline") or "").strip()
    if changed_text:
        try:
            changed_deadline = date.fromisoformat(changed_text)
        except ValueError as error:
            raise ValueError("Enter a valid changed deadline.") from error

    increase_text = str(form.get("budget_increase") or "").strip()
    if increase_text:
        try:
            budget_increase = float(increase_text.replace(",", ""))
        except ValueError as error:
            raise ValueError("Budget increase must be a number.") from error
        if not math.isfinite(budget_increase) or budget_increase < 0 or budget_increase > 1_000_000_000_000:
            raise ValueError("Budget increase must be between 0 and 1,000,000,000,000.")
    else:
        budget_increase = 0.0
    if budget_increase and project.get("budget") is None:
        raise ValueError("Record a project budget before testing a budget increase.")

    current_deadline = (
        date.fromisoformat(project["expected_completion_date"])
        if project.get("expected_completion_date") else None
    )
    scenario_deadline = changed_deadline or current_deadline
    if delay_days and scenario_deadline is None:
        raise ValueError("Add a current or changed deadline before testing extra delay days.")
    if scenario_deadline and delay_days:
        try:
            scenario_deadline += timedelta(days=delay_days)
        except OverflowError as error:
            raise ValueError("That delay would move the deadline outside the supported date range.") from error
    start_text = project.get("start_date")
    if scenario_deadline and start_text and scenario_deadline < date.fromisoformat(start_text):
        raise ValueError("The scenario deadline cannot be before the project start date.")

    hypothetical = dict(project)
    if scenario_deadline:
        hypothetical["expected_completion_date"] = scenario_deadline.isoformat()
    if project.get("budget") is not None:
        hypothetical["budget"] = float(project["budget"]) + budget_increase

    expenditure = project.get("expenditure")
    original_budget = project.get("budget")
    scenario_budget = hypothetical.get("budget")

    def utilization(amount):
        if expenditure is None or amount is None or amount <= 0:
            return None
        return expenditure / amount * 100

    return {
        "baseline_risk": _project_delay_assessment(project, today),
        "scenario_risk": _project_delay_assessment(hypothetical, today),
        "current_deadline": current_deadline.isoformat() if current_deadline else None,
        "scenario_deadline": scenario_deadline.isoformat() if scenario_deadline else None,
        "current_budget": original_budget,
        "scenario_budget": scenario_budget,
        "current_utilization": utilization(original_budget),
        "scenario_utilization": utilization(scenario_budget),
        "delay_days": delay_days,
        "budget_increase": budget_increase,
    }


@app.route("/", methods=["GET", "POST"])
def dashboard():
    projects = get_projects()
    today = date.today()
    selected_id = request.args.get("project", type=int)
    simulation = None
    simulation_error = None
    simulation_form = {"delay_days": "", "changed_deadline": "", "budget_increase": ""}
    if request.method == "POST":
        if request.form.get("action") != "simulate":
            abort(400)
        posted_id = request.form.get("record_id", type=int)
        if not posted_id or (selected_id and selected_id != posted_id):
            abort(400)
        selected_id = posted_id
        simulation_form = {
            "delay_days": request.form.get("delay_days", ""),
            "changed_deadline": request.form.get("changed_deadline", ""),
            "budget_increase": request.form.get("budget_increase", ""),
        }
    query = request.args.get("q", "").strip().lower()
    status_filter = request.args.get("status", "").strip()
    visible_projects = [
        project for project in projects
        if (not query or query in " ".join(str(project.get(key) or "") for key in ("project_id", "name", "project_type", "location")).lower())
        and (not status_filter or project["status"] == status_filter)
    ]
    for project in visible_projects:
        project["warnings"] = _project_warnings(project, today)
    for project in projects:
        project["warnings"] = _project_warnings(project, today)
    counts = {
        "total": len(projects),
        "in_progress": sum(project["status"] == "In Progress" for project in projects),
        "completed": sum(project["status"] == "Completed" for project in projects),
        "warnings": sum(bool(project["warnings"]) for project in projects),
    }
    budget_total = sum(project["budget"] or 0 for project in projects)
    expenditure_total = sum(project["expenditure"] or 0 for project in projects)
    budget_count = sum(project["budget"] is not None for project in projects)
    expenditure_count = sum(project["expenditure"] is not None for project in projects)
    deadline_list = sorted(
        (
            project for project in projects
            if project["status"] != "Completed" and project.get("expected_completion_date")
            and project["expected_completion_date"] >= today.isoformat()
        ),
        key=lambda project: project["expected_completion_date"],
    )[:5]
    status_counts = {
        status: sum(project["status"] == status for project in projects)
        for status in PROJECT_STATUSES
    }
    selected = None
    edit_project = None
    if selected_id:
        selected = get_project(selected_id)
        if selected is None:
            abort(404)
        selected["warnings"] = _project_warnings(selected, today)
        if request.args.get("edit") == "1":
            edit_project = selected
        if request.method == "POST":
            try:
                simulation = _simulate_project(selected, request.form, today)
            except ValueError as error:
                simulation_error = str(error)
        elif request.args.get("assess") == "1":
            selected["delay_assessment"] = _project_delay_assessment(selected, today)
    preview = None
    token = request.args.get("import")
    if token:
        with connect() as connection:
            row = connection.execute("SELECT rows_json FROM import_batches WHERE token = ?", (token,)).fetchone()
        if row:
            preview = {"token": token, "rows": json.loads(row["rows_json"])}
    return render_template(
        "dashboard.html", projects=visible_projects, counts=counts,
        budget_total=budget_total, expenditure_total=expenditure_total,
        budget_count=budget_count, expenditure_count=expenditure_count,
        deadlines=deadline_list, status_counts=status_counts,
        statuses=PROJECT_STATUSES, selected=selected, edit_project=edit_project,
        simulation=simulation, simulation_error=simulation_error,
        simulation_form=simulation_form,
        preview=preview, query=request.args.get("q", ""),
        status_filter=status_filter,
    )


@app.route("/projects", methods=["POST"])
def project_save():
    record_id = request.form.get("record_id", type=int)
    existing = get_project(record_id) if record_id else None
    if record_id and existing is None:
        abort(404)
    project, errors = _validate_project(request.form)
    if not errors:
        with connect() as connection:
            duplicate = connection.execute(
                "SELECT id FROM projects WHERE project_id = ? COLLATE NOCASE AND id != ?",
                (project["project_id"], record_id or -1),
            ).fetchone()
        if duplicate:
            errors.append("That project ID is already in use.")
    if errors:
        for error in errors:
            flash(error, "error")
        return redirect(url_for("dashboard", project=record_id) if record_id else url_for("dashboard"))
    try:
        save_project(project, record_id)
    except sqlite3.IntegrityError:
        flash("That project ID is already in use.", "error")
        return redirect(url_for("dashboard", project=record_id) if record_id else url_for("dashboard"))
    flash("Project changes saved." if record_id else "Project added.", "success")
    return redirect(url_for("dashboard", project=record_id) if record_id else url_for("dashboard"))


@app.route("/projects/<int:record_id>/delete", methods=["POST"])
def project_delete(record_id):
    if get_project(record_id) is None:
        abort(404)
    delete_project(record_id)
    flash("Project deleted.", "success")
    return redirect(url_for("dashboard"))


@app.route("/projects/template.csv", methods=["GET"])
def import_template():
    headers = list(IMPORT_COLUMNS.keys())
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    return app.response_class(
        output.getvalue(), mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=uptrack_import_template.csv"},
    )


@app.route("/imports/preview", methods=["POST"])
def import_preview():
    upload = request.files.get("file")
    if not upload or not upload.filename:
        flash("Choose a CSV or .xlsx file to preview.", "error")
        return redirect(url_for("dashboard"))
    try:
        raw_rows = _read_import(upload)
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("dashboard"))
    seen_ids = set()
    with connect() as connection:
        existing_ids = {
            row["project_id"].casefold()
            for row in connection.execute("SELECT project_id FROM projects").fetchall()
        }
    preview_rows = []
    for number, raw in enumerate(raw_rows, 2):
        project, errors = _validate_project(raw)
        normalized_id = project["project_id"].casefold()
        if normalized_id and (normalized_id in existing_ids or normalized_id in seen_ids):
            errors.append("Project ID duplicates an existing project or another row in this file.")
        if normalized_id:
            seen_ids.add(normalized_id)
        preview_rows.append({"row": number, "project": project, "errors": errors})
    token = uuid.uuid4().hex
    with connect() as connection:
        connection.execute("DELETE FROM import_batches WHERE created_at < datetime('now', '-1 day')")
        connection.execute(
            "INSERT INTO import_batches (token, rows_json) VALUES (?, ?)",
            (token, json.dumps(preview_rows, ensure_ascii=False)),
        )
        connection.commit()
    return redirect(url_for("dashboard", **{"import": token}))


@app.route("/imports/commit", methods=["POST"])
def import_commit():
    token = request.form.get("token", "")
    with connect() as connection:
        batch = connection.execute("SELECT rows_json FROM import_batches WHERE token = ?", (token,)).fetchone()
        if batch is None:
            flash("That import preview has expired. Upload the file again.", "error")
            return redirect(url_for("dashboard"))
        rows = json.loads(batch["rows_json"])
        errors = [f"Row {row['row']}: {error}" for row in rows for error in row["errors"]]
        current_ids = {
            row["project_id"].casefold()
            for row in connection.execute("SELECT project_id FROM projects").fetchall()
        }
        incoming_ids = [row["project"]["project_id"].casefold() for row in rows]
        if any(not value or value in current_ids for value in incoming_ids):
            errors.append("A project ID now conflicts with a saved project. No rows were imported.")
        if len(incoming_ids) != len(set(incoming_ids)):
            errors.append("Duplicate project IDs were found. No rows were imported.")
        if errors:
            for error in errors[:30]:
                flash(error, "error")
            return redirect(url_for("dashboard", **{"import": token}))
        fields = tuple(rows[0]["project"].keys()) if rows else ()
        if not fields:
            flash("The import contains no rows.", "error")
            return redirect(url_for("dashboard"))
        placeholders = ", ".join("?" for _ in fields)
        statement = f"INSERT INTO projects ({', '.join(fields)}) VALUES ({placeholders})"
        try:
            connection.executemany(
                statement,
                [
                    tuple(json.dumps(row["project"][field], ensure_ascii=False) if field == "milestones" else row["project"][field] for field in fields)
                    for row in rows
                ],
            )
        except sqlite3.IntegrityError:
            connection.rollback()
            flash("A project ID conflicts with a saved record. No rows were imported.", "error")
            return redirect(url_for("dashboard", **{"import": token}))
        connection.execute("DELETE FROM import_batches WHERE token = ?", (token,))
        connection.commit()
    flash(f"Imported {len(rows)} projects.", "success")
    return redirect(url_for("dashboard"))


@app.errorhandler(413)
def too_large(_error):
    flash("The file is too large. Maximum upload size is 3 MB.", "error")
    return redirect(url_for("dashboard"))


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)